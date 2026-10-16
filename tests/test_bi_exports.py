"""The files the BI tools read (bi/data/*.csv) and the Excel model
(bi/excel/AP_business_case.xlsx) must match what the code computes today.
If a test fails after a code or data change, rerun
`python scripts/export_bi_data.py` and
`python scripts/build_excel_workbook.py bi/excel/AP_business_case.xlsx`.
"""

import importlib.util
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src import business_case

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "bi" / "excel" / "AP_business_case.xlsx"


def load_script(name: str):
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def exporter():
    return load_script("export_bi_data")


@pytest.fixture(scope="module")
def exports(exporter):
    return exporter.build_exports()


def test_committed_exports_are_up_to_date(exporter, exports):
    for name, frame in exports.items():
        committed = (exporter.BI_DATA_DIR / f"{name}.csv").read_text(encoding="utf-8")
        assert exporter.to_csv_text(frame) == committed, f"bi/data/{name}.csv is stale"


def test_no_stray_export_files(exporter, exports):
    on_disk = {p.stem for p in exporter.BI_DATA_DIR.glob("*.csv")}
    assert on_disk == set(exports)


def test_headline_figures_match_the_app(exports):
    figures = exports["headline_figures"].set_index("figure")["value"]
    case = business_case.compute_business_case()
    assert figures["npv_hard_eur"] == case["base"]["hard"]["npv_eur"]
    assert figures["hard_leakage_eur"] == pytest.approx(73_552.87)
    assert exports["invoices"]["amount"].sum() == pytest.approx(figures["total_spend_eur"], abs=0.5)


def test_workbook_checks_carry_todays_python_figures(tmp_path):
    """The Checks sheet compares the live formulas with the Python figures written in at build time.
    A fresh build must write the same expected values as the committed workbook."""
    fresh = tmp_path / "fresh.xlsx"
    load_script("build_excel_workbook").build(fresh)

    def expected(path):
        ws = load_workbook(path)["Checks"]
        return [(r[0], r[2]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[2] is not None]

    assert expected(WORKBOOK) == expected(fresh)
    base = business_case.compute_business_case()["base"]
    rows = dict(expected(WORKBOOK))
    assert rows["NPV, hard benefits"] == base["hard"]["npv_eur"]


def test_workbook_uses_formulas_excel_and_libreoffice_both_evaluate():
    wb = load_workbook(WORKBOOK)
    formulas = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) > 400
    banned = ["XLOOKUP", "XMATCH", "FILTER(", "SORT(", "UNIQUE(", "SEQUENCE("]
    assert not [f for f in formulas if any(b in f.upper() for b in banned)]
    assert wb["Inputs"]["B1"].value == "Base"
